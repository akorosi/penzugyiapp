"""CIB folyószámlakivonat (.xls/.xlsx) beolvasása.

Elvárt szerkezet:
 - A11-től kezdve a konkrét tranzakciók, időben visszafelé (legújabb elöl).
 - A oszlop: dátum
 - B oszlop: tranzakció típusa
 - C oszlop: közlemény
 - D oszlop: összeg (kiadás = negatív, bevétel = pozitív)

A fájl A11-től kezdődő teljes tartalma beolvasásra kerül, dátum/hónap
szerinti szűrés NINCS. A fájlt a memóriából olvassuk (Lambda-ban nincs
tartós lemez), pandas nélkül — csak xlrd (.xls) és openpyxl (.xlsx).
"""

from __future__ import annotations

import hashlib
import io
import re
from datetime import date, datetime, timedelta

from dateutil import parser as date_parser

FIRST_DATA_ROW = 10  # 0-indexelt: A11 az első adatsor
_EXCEL_EPOCH = datetime(1899, 12, 30)


def _read_rows_xls(content: bytes) -> list[list]:
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    sheet = book.sheet_by_index(0)
    rows = []
    for r in range(FIRST_DATA_ROW, sheet.nrows):
        row = []
        for c in range(min(4, sheet.ncols)):
            cell = sheet.cell(r, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
            elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                row.append(None)
            else:
                row.append(cell.value)
        rows.append(row)
    return rows


def _read_rows_xlsx(content: bytes) -> list[list]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = [list(r[:4]) for r in ws.iter_rows(min_row=FIRST_DATA_ROW + 1, values_only=True)]
    wb.close()
    return rows


def _to_date(value) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Excel dátum-sorszám
        if 1 <= value < 2958466:
            return (_EXCEL_EPOCH + timedelta(days=float(value))).date()
        return None
    text = str(value).strip()
    if not text:
        return None
    # "2024. 03. 15." / "2024.03.15" / "2024-03-15"
    text = re.sub(r"\s+", "", text).rstrip(".")
    try:
        return date_parser.parse(text, yearfirst=True, dayfirst=False).date()
    except (ValueError, OverflowError):
        return None


def _to_amount(value) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("−", "-")
    text = re.sub(r"[\s ]|Ft|HUF", "", text, flags=re.IGNORECASE)
    if "," in text and "." not in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def _to_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_cib_statement(content: bytes) -> list[dict]:
    rows = None
    last_err: Exception | None = None
    for reader in (_read_rows_xls, _read_rows_xlsx):
        try:
            rows = reader(content)
            break
        except Exception as e:  # próbáljuk a következő formátumot
            last_err = e
    if rows is None:
        raise ValueError(f"Nem sikerült beolvasni az Excel fájlt: {last_err}")

    records: list[dict] = []
    for row in rows:
        row = list(row) + [None] * (4 - len(row))
        d = _to_date(row[0])
        if d is None:
            continue
        amount = _to_amount(row[3])
        if amount is None:
            continue
        records.append(
            {
                "date": d,
                "tx_type": _to_text(row[1]),
                "description": _to_text(row[2]),
                "amount": amount,
            }
        )
    return records


def make_hash(rec: dict, owner: str) -> str:
    """Duplikátum-azonosító. A tulajdonos is része, így ugyanazt a kivonatot két
    felhasználó egymástól függetlenül feltöltheti."""
    key = f"{owner}|{rec['date'].isoformat()}|{rec['tx_type']}|{rec['description']}|{rec['amount']:.2f}"
    return hashlib.sha256(key.encode("utf-8")).hexdigest()
