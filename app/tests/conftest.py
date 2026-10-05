"""Teszt-beállítás: valódi PostgreSQL ellen futnak (TEST_DATABASE_URL), mert a
DSQL-kompatibilis SQL-t nem lehet SQLite-tal hitelesen tesztelni.

    TEST_DATABASE_URL=postgresql://postgres@localhost:5432/penzugy_test pytest

A hitelesítés Cognito módban fut; a Cognito/Google oldalt a tesztek helyettesítik.
"""

import io
import os
import sys
from datetime import date

import pytest

TEST_DB = os.environ.get("TEST_DATABASE_URL")
if TEST_DB:
    os.environ["DATABASE_URL"] = TEST_DB
os.environ.pop("DSQL_ENDPOINT", None)
os.environ.pop("AWS_LAMBDA_FUNCTION_NAME", None)
os.environ.update(
    AUTH_MODE="cognito",
    SESSION_SECRET="test-session-secret-0123456789abcdef0123456789",
    ALLOWED_EMAILS="anna@example.com, Bela@Example.com",
    COGNITO_CLIENT_ID="test-client",
    COGNITO_CLIENT_SECRET="test-secret",
    APP_URL="https://app.example.net",
    COGNITO_DOMAIN="https://penzugyek.auth.eu-central-1.amazoncognito.com",
    COGNITO_USER_POOL_ID="eu-central-1_TEST",
    COGNITO_REGION="eu-central-1",
)

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

USER_A = "anna@example.com"
USER_B = "bela@example.com"
# A módosító kérésekhez szükséges CSRF fejléc (a frontend mindig küldi).
HDR = {"X-Requested-With": "penzugyek"}


def pytest_collection_modifyitems(config, items):
    if TEST_DB:
        return
    skip = pytest.mark.skip(reason="TEST_DATABASE_URL nincs beállítva")
    for item in items:
        if "needs_db" in item.keywords:
            item.add_marker(skip)


def pytest_configure(config):
    config.addinivalue_line("markers", "needs_db: PostgreSQL adatbázist igénylő teszt")


def login(client, email: str) -> None:
    import auth

    client.set_cookie(auth.SESSION_COOKIE, auth.make_session(email))


@pytest.fixture()
def app_client():
    """Üres adatbázis, bejelentkezés nélküli kliens."""
    import db
    from main import app

    db.migrate()
    conn = db.get_conn()
    conn.execute("DELETE FROM attributes")
    conn.execute("DELETE FROM transactions")
    conn.execute("DELETE FROM category_rules")
    return app.test_client()


@pytest.fixture()
def client(app_client):
    """USER_A-ként bejelentkezett kliens."""
    login(app_client, USER_A)
    return app_client


def make_xlsx(rows: list[tuple]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 1, "CIB Bank — számlakivonat")
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            ws.cell(11 + i, 1 + j, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


SAMPLE_ROWS = [
    (date(2026, 9, 28), "Kártyás vásárlás", "SPAR 123 BUDAPEST", -12345),
    (date(2026, 9, 27), "Kártyás vásárlás", "MAGYAR ALLAMKINCSTAR", -50000),
    (date(2026, 9, 26), "Jóváírás", "MUNKABER ACME KFT", 650000),
    ("2026.09.25", "Kártyás vásárlás", "Ismeretlen bolt", "-1 500"),
    (None, None, "összesen sor, dátum nélkül", 0),
]
